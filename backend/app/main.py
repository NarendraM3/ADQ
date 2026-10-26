import json, logging, mimetypes, re
import asyncio
from contextlib import asynccontextmanager
from pathlib import Path
from fastapi import BackgroundTasks, Depends, FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse
from io import BytesIO
import csv
from cryptography.fernet import Fernet, InvalidToken
from sqlalchemy import delete, or_, select
from sqlalchemy.orm import Session
from app.config import settings
from app.database import Base, engine, get_db, SessionLocal
from app.models import Correction, Credential, Dataset, Execution, Job, Rule, WorkflowSelection
from app.processing import destination_for_job, encode_csv, execute, parse_data
from app.rules import RULE_TYPES, rule_signature, validate_rule_definition
from app.connectors.destinations import read_databricks_table, write_destination
from app.schemas import CorrectionCreate, CredentialCreate, JobConfiguration, JobCreate, JobUpdate, RuleCreate, RuleUpdate, normalize_platform
from app.connectors.storage import LocalStorage, get_storage, get_source_adapter

logging.basicConfig(level=settings.log_level)
@asynccontextmanager
async def lifespan(app):
    Base.metadata.create_all(bind=engine)
    # create_all does not add columns to an existing SQLite database.
    if engine.dialect.name == "sqlite":
        from sqlalchemy import inspect
        if "settings" not in {column["name"] for column in inspect(engine).get_columns("jobs")}:
            with engine.begin() as connection:
                connection.exec_driver_sql("ALTER TABLE jobs ADD COLUMN settings JSON NOT NULL DEFAULT '{}'")
        if "environment" not in {column["name"] for column in inspect(engine).get_columns("jobs")}:
            with engine.begin() as connection:
                # Existing jobs have no reliable record of the platform selected at creation.
                connection.exec_driver_sql("ALTER TABLE jobs ADD COLUMN environment VARCHAR(100)")
        if "rule_snapshot" not in {column["name"] for column in inspect(engine).get_columns("executions")}:
            with engine.begin() as connection:
                connection.exec_driver_sql("ALTER TABLE executions ADD COLUMN rule_snapshot JSON")
    # BackgroundTasks are process-local. Resume persisted queued executions
    # after a restart; execute() atomically claims each row before doing work.
    with SessionLocal() as db:
        queued_ids = list(db.scalars(select(Execution.id).where(Execution.status == "queued")))
    for execution_id in queued_ids:
        asyncio.create_task(asyncio.to_thread(execute, execution_id))
    yield


app = FastAPI(title="ADQ Automation API", version="0.1.0", description="AI-assisted data quality jobs. Local CSV/JSON processing is ready; cloud providers require credentials and optional SDKs.", lifespan=lifespan)


def not_found(obj, name):
    if obj is None: raise HTTPException(404, f"{name} not found")
    return obj


def job_environment(job: Job) -> str | None:
    """Return this job's saved platform, with explicit legacy metadata as fallback."""
    settings = job.settings if isinstance(job.settings, dict) else {}
    destination = job.destination if isinstance(job.destination, dict) else {}
    return normalize_platform(
        job.environment
        or settings.get("environment")
        or settings.get("platform")
        or destination.get("platform")
    )


def infer_job_environment(environment: str | None, settings: dict | None = None) -> str | None:
    settings = settings if isinstance(settings, dict) else {}
    # The selected platform is independent of the job's source and output destination.
    return normalize_platform(environment or settings.get("environment") or settings.get("platform"))


def job_response(job: Job) -> dict:
    return {
        "id": job.id, "name": job.name, "description": job.description,
        "dataset_ids": job.dataset_ids, "selected_columns": job.selected_columns,
        "join_config": job.join_config, "destination": job.destination,
        "environment": job_environment(job), "settings": job.settings,
        "created_at": job.created_at, "updated_at": job.updated_at,
    }


def credential_key():
    key = settings.secret_encryption_key
    if not key:
        raise HTTPException(503, "Credential storage is unavailable: SECRET_ENCRYPTION_KEY is missing or empty. Set it in the project .env or backend environment.")
    try:
        return Fernet(key.strip().encode("ascii"))
    except (UnicodeEncodeError, TypeError, ValueError) as e:
        raise HTTPException(503, "Credential storage is unavailable: SECRET_ENCRYPTION_KEY is not a valid Fernet key. Check its format and whether the backend environment overrides the project .env.") from e


def sensitive_config_keys(value, prefix=""):
    found=[]
    if isinstance(value,dict):
        for key,item in value.items():
            name=f"{prefix}.{key}" if prefix else str(key)
            if any(word in str(key).lower() for word in ("secret","password","token","connection_string","private_key")):
                found.append(name)
            else: found.extend(sensitive_config_keys(item,name))
    elif isinstance(value,list):
        for index,item in enumerate(value): found.extend(sensitive_config_keys(item,f"{prefix}[{index}]"))
    return found


@app.get("/health")
def health(): return {"status": "ok", "database": "configured", "local_workflow": "ready", "ai": settings.ai_provider if settings.ai_provider != "none" else "disabled"}


@app.get("/api/v1/capabilities")
def capabilities():
    return {"sources": {"local_csv_json": "ready", "s3": "optional boto3 adapter", "azure_blob": "optional Azure SDK adapter", "onelake": "optional Fabric OAuth adapter", "databricks": "optional SQL connector with table discovery", "snowflake": "optional SQL connector with table discovery"},
        "destinations": {"local": "ready", "s3": "optional boto3 adapter", "azure_blob": "optional Azure SDK adapter", "onelake": "optional Fabric OAuth adapter", "databricks": "optional SQL connector", "snowflake": "optional SQL connector", "fabric": "optional OneLake adapter"},
        "ai_suggestions": "configured" if settings.ai_provider != "none" and settings.resolved_ai_api_key and settings.resolved_ai_model and settings.resolved_ai_base_url else "disabled; set AI_PROVIDER, AI_API_KEY and AI_MODEL (or GROQ_API_KEY and GROQ_MODEL for Groq)"}


@app.get("/api/v1/datasets")
def list_datasets(db: Session = Depends(get_db)):
    return db.scalars(select(Dataset).order_by(Dataset.id.desc())).all()


@app.get("/api/v1/workflows/{workflow_id}/selection")
def get_workflow_selection(workflow_id: str, db: Session = Depends(get_db)):
    selection = db.get(WorkflowSelection, workflow_id)
    if selection is None:
        return {"workflow_id": workflow_id, "dataset_ids": [], "source_connection_ids": {}, "datasets": []}
    ids = selection.dataset_ids or []
    rows = {d.id: d for d in db.scalars(select(Dataset).where(Dataset.id.in_(ids))).all()} if ids else {}
    valid = [did for did in ids if did in rows]
    return {"workflow_id": workflow_id, "dataset_ids": valid,
            "source_connection_ids": selection.source_connection_ids or {},
            "datasets": [rows[did] for did in valid]}


@app.put("/api/v1/workflows/{workflow_id}/selection")
def set_workflow_selection(workflow_id: str, body: dict, db: Session = Depends(get_db)):
    """Persist exactly the explicitly selected dataset IDs for this workflow."""
    ids = body.get("dataset_ids")
    connections = body.get("source_connection_ids", {})
    if not isinstance(workflow_id, str) or not workflow_id.strip() or len(workflow_id) > 36:
        raise HTTPException(422, "A valid workflow ID is required")
    if not isinstance(ids, list) or any(not isinstance(i, int) or i <= 0 for i in ids):
        raise HTTPException(422, "dataset_ids must contain valid dataset IDs")
    if len(set(ids)) != len(ids):
        raise HTTPException(422, "dataset_ids must be unique")
    rows = {d.id: d for d in db.scalars(select(Dataset).where(Dataset.id.in_(ids))).all()}
    if len(rows) != len(ids):
        raise HTTPException(404, "A selected dataset no longer exists")
    normalized_connections = {}
    for dataset_id in ids:
        expected = connections.get(str(dataset_id)) if isinstance(connections, dict) else None
        source = (rows[dataset_id].schema_json or {}).get("_source", {})
        actual = source.get("credential_id")
        if expected is not None and actual != expected:
            raise HTTPException(422, "Selected dataset does not belong to the selected source connection")
        if actual is not None:
            normalized_connections[str(dataset_id)] = actual
    selection = db.get(WorkflowSelection, workflow_id)
    if selection is None:
        selection = WorkflowSelection(workflow_id=workflow_id, dataset_ids=ids,
                                      source_connection_ids=normalized_connections)
        db.add(selection)
    else:
        selection.dataset_ids = ids
        selection.source_connection_ids = normalized_connections
    db.commit()
    return {"workflow_id": workflow_id, "dataset_ids": ids,
            "source_connection_ids": normalized_connections,
            "datasets": [rows[dataset_id] for dataset_id in ids]}


@app.post("/api/v1/datasets/upload", status_code=201)
async def upload_dataset(file: UploadFile = File(...), name: str | None = None, db: Session = Depends(get_db)):
    fname = Path(file.filename or "").name
    if Path(fname).suffix.lower() not in (".csv", ".tsv", ".json", ".xlsx", ".xls", ".parquet"): raise HTTPException(415, "Upload a CSV, TSV, JSON, Excel, or Parquet file")
    content = await file.read(settings.max_upload_mb * 1024 * 1024 + 1)
    if len(content) > settings.max_upload_mb * 1024 * 1024: raise HTTPException(413, "Upload exceeds MAX_UPLOAD_MB")
    try: records = parse_data(fname, content)
    except Exception as e: raise HTTPException(400, f"Invalid dataset: {e}") from e
    columns = list(dict.fromkeys(k for row in records[:1000] for k in row))
    schema = {col: {"inferred_type": infer_type([r.get(col) for r in records[:1000]]), "nullable": any(r.get(col) in (None, "") for r in records[:1000])} for col in columns}
    path = LocalStorage().save(f"uploads/{__import__('uuid').uuid4().hex}-{fname}", content)
    dataset = Dataset(name=name or Path(fname).stem, filename=fname, path=path, schema_json={"columns": schema, "row_count": len(records)})
    db.add(dataset); db.commit(); db.refresh(dataset); return dataset


def infer_type(values):
    vals = [str(v) for v in values if v not in (None, "")]
    if vals and all(is_int(v) for v in vals): return "integer"
    if vals and all(is_float(v) for v in vals): return "number"
    if vals and all("@" in v for v in vals): return "email-like"
    return "string"


def is_int(v):
    try: int(v); return True
    except (ValueError, TypeError): return False


def is_float(v):
    try: float(v); return True
    except (ValueError, TypeError): return False


@app.get("/api/v1/datasets/{dataset_id}")
def dataset_detail(dataset_id: int, db: Session = Depends(get_db)):
    return not_found(db.get(Dataset, dataset_id), "Dataset")


@app.get("/api/v1/jobs")
def list_jobs(db: Session = Depends(get_db)):
    return [job_response(job) for job in db.scalars(select(Job).order_by(Job.id.desc()))]


@app.post("/api/v1/jobs", status_code=201)
def create_job(body: JobCreate, db: Session = Depends(get_db)):
    datasets=[db.get(Dataset,did) for did in body.dataset_ids]
    if any(d is None for d in datasets): raise HTTPException(422, "Every dataset_ids entry must exist")
    if "selected_columns" in body.model_fields_set and not body.selected_columns:
        raise HTTPException(422,"Select at least one column for the job")
    validate_job_inputs(body,datasets)
    values = body.model_dump()
    values["environment"] = infer_job_environment(values.get("environment"), body.settings)
    job = Job(**values); db.add(job); db.commit(); db.refresh(job); return job_response(job)


@app.get("/api/v1/jobs/{job_id}")
def get_job(job_id: int, db: Session = Depends(get_db)):
    return job_response(not_found(db.get(Job, job_id), "Job"))


@app.patch("/api/v1/jobs/{job_id}")
def update_job(job_id: int, body: JobUpdate, db: Session = Depends(get_db)):
    job = not_found(db.get(Job, job_id), "Job")
    changes=body.model_dump(exclude_unset=True)
    dataset_ids=changes.get("dataset_ids",job.dataset_ids)
    datasets=[db.get(Dataset,did) for did in dataset_ids]
    if any(d is None for d in datasets): raise HTTPException(422, "Every dataset_ids entry must exist")
    merged={"name":changes.get("name",job.name),"description":changes.get("description",job.description),"dataset_ids":dataset_ids,
        "selected_columns":changes.get("selected_columns",job.selected_columns),"join_config":changes.get("join_config",job.join_config),
        "destination":changes.get("destination",job.destination),"environment":changes.get("environment",job_environment(job)),"settings":changes.get("settings",job.settings or {})}
    validate_job_inputs(JobCreate(**merged),datasets)
    for k,v in changes.items(): setattr(job, k, v)
    db.commit(); db.refresh(job); return job_response(job)


def validate_configuration_rules(body: JobCreate, rules: list[RuleCreate], db: Session):
    for rule in rules:
        kind = rule.rule_type.upper()
        try:
            validate_rule_definition(kind, rule.parameters)
        except ValueError as e:
            raise HTTPException(422, str(e)) from e
        if rule.dataset_id is not None:
            if rule.dataset_id not in body.dataset_ids:
                raise HTTPException(422, "Rule dataset_id must be selected by the job")
            ds = db.get(Dataset, rule.dataset_id)
            if not ds or rule.column_name not in ds.schema_json.get("columns", {}):
                raise HTTPException(422, "Rule column does not exist in the selected dataset")


@app.post("/api/v1/jobs/configuration", status_code=201)
def create_job_configuration(body: JobConfiguration, db: Session = Depends(get_db)):
    datasets = [db.get(Dataset, did) for did in body.job.dataset_ids]
    if any(d is None for d in datasets):
        raise HTTPException(422, "Every dataset_ids entry must exist")
    if "selected_columns" in body.job.model_fields_set and not body.job.selected_columns:
        raise HTTPException(422,"Select at least one column for the job")
    validate_job_inputs(body.job, datasets)
    validate_configuration_rules(body.job, body.rules, db)
    values = body.job.model_dump()
    values["environment"] = infer_job_environment(values.get("environment"), body.job.settings)
    job = Job(**values)
    db.add(job)
    db.flush()
    for rule in body.rules:
        values = rule.model_dump()
        values["rule_type"] = values["rule_type"].upper()
        values["enabled"] = bool(values["enabled"])
        values["approved"] = bool(values["approved"])
        db.add(Rule(job_id=job.id, **values))
    db.commit()
    db.refresh(job)
    return job_response(job)


@app.put("/api/v1/jobs/{job_id}/configuration")
def replace_job_configuration(job_id: int, body: JobConfiguration, db: Session = Depends(get_db)):
    job = not_found(db.get(Job, job_id), "Job")
    datasets = [db.get(Dataset, did) for did in body.job.dataset_ids]
    if any(d is None for d in datasets):
        raise HTTPException(422, "Every dataset_ids entry must exist")
    validate_job_inputs(body.job, datasets)
    validate_configuration_rules(body.job, body.rules, db)
    values = body.job.model_dump()
    values["environment"] = infer_job_environment(values.get("environment") or job_environment(job), body.job.settings)
    for key, value in values.items():
        setattr(job, key, value)
    db.execute(delete(Rule).where(Rule.job_id == job_id))
    for rule in body.rules:
        values = rule.model_dump()
        values["rule_type"] = values["rule_type"].upper()
        values["enabled"] = bool(values["enabled"])
        values["approved"] = bool(values["approved"])
        db.add(Rule(job_id=job_id, **values))
    db.commit()
    db.refresh(job)
    return job_response(job)


@app.delete("/api/v1/jobs/{job_id}", status_code=204)
def delete_job(job_id: int, db: Session = Depends(get_db)):
    job = not_found(db.get(Job, job_id), "Job")
    # Explicitly remove dependent records so deletion is safe with SQLite and
    # databases where foreign-key cascade enforcement is not enabled.
    executions = list(db.scalars(select(Execution).where(Execution.job_id == job_id)))
    execution_ids = [item.id for item in executions]
    if execution_ids:
        db.execute(delete(Correction).where(Correction.execution_id.in_(execution_ids)))
        db.execute(delete(Execution).where(Execution.job_id == job_id))
    db.execute(delete(Rule).where(Rule.job_id == job_id))
    db.delete(job)
    db.commit()


def validate_job_inputs(body,datasets):
    if len(set(body.dataset_ids)) != len(body.dataset_ids): raise HTTPException(422,"dataset_ids must be unique")
    if len(datasets)>2 and (body.settings or {}).get("combine_mode") != "union": raise HTTPException(422,"A join job can combine up to two datasets")
    known={str(d.id):set((d.schema_json or {}).get("columns",{})) for d in datasets}
    for dataset_id,columns in body.selected_columns.items():
        if str(dataset_id) not in known: raise HTTPException(422,f"selected_columns references dataset {dataset_id} outside the job")
        missing=set(columns)-known[str(dataset_id)]
        if missing: raise HTTPException(422,f"Unknown selected columns for dataset {dataset_id}: {sorted(missing)}")
    if body.selected_columns and not any(body.selected_columns.values()):
        raise HTTPException(422,"Select at least one column for the job")
    if len(datasets)>=2 and (body.settings or {}).get("combine_mode") != "union":
        cfg=body.join_config
        if not cfg: raise HTTPException(422,"join_config is required for a two-dataset job")
        try: left_id,right_id=int(cfg.get("left_dataset_id",-1)),int(cfg.get("right_dataset_id",-1))
        except (TypeError,ValueError): raise HTTPException(422,"join_config requires integer dataset IDs")
        if left_id not in [d.id for d in datasets] or right_id not in [d.id for d in datasets]:
            raise HTTPException(422,"join_config dataset IDs must be selected by the job")
        if left_id==right_id: raise HTTPException(422,"join_config must reference two different datasets")
        for did,key in ((left_id,cfg.get("left_key")),(right_id,cfg.get("right_key"))):
            if key not in known[str(did)]: raise HTTPException(422,f"Join key {key!r} does not exist in dataset {did}")


@app.get("/api/v1/jobs/{job_id}/landing-zone")
def landing_zone(job_id: int, db: Session = Depends(get_db)):
    job = not_found(db.get(Job, job_id), "Job")
    return {"job_id": job_id, "files": [d for d in db.scalars(select(Dataset).where(Dataset.id.in_(job.dataset_ids))).all()], "selected_columns": job.selected_columns}


@app.get("/api/v1/jobs/{job_id}/rules")
def list_rules(job_id: int, db: Session = Depends(get_db)):
    not_found(db.get(Job, job_id), "Job")
    return db.scalars(select(Rule).where(Rule.job_id == job_id).order_by(Rule.id)).all()


RULE_IMPORT_COLUMNS = ["dataset_id", "column_name", "rule_type", "parameters", "severity", "enabled", "approved"]


def _rules_workbook():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    wb = Workbook(); ws = wb.active; ws.title = "Rules"
    ws.append(RULE_IMPORT_COLUMNS)
    ws.append(["", "email", "EMAIL_FORMAT", "{}", "high", True, False])
    ws.append(["", "age", "RANGE", '{"min":0,"max":120}', "medium", True, False])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = "A1:G3"
    widths = [16, 26, 20, 38, 14, 12, 12]
    for i, width in enumerate(widths, 1): ws.column_dimensions[chr(64+i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="173B52"); cell.alignment = Alignment(vertical="center")
    dv = DataValidation(type="list", formula1='"' + ",".join(sorted(RULE_TYPES)) + '"'); ws.add_data_validation(dv); dv.add("C2:C1000")
    sev = DataValidation(type="list", formula1='"low,medium,high,critical"'); ws.add_data_validation(sev); sev.add("E2:E1000")
    flags = DataValidation(type="list", formula1='"TRUE,FALSE"'); ws.add_data_validation(flags); flags.add("F2:G1000")
    info=wb.create_sheet("Instructions")
    for row in [["AI Data Quality Rules — Import Guide"], ["Required columns: column_name, rule_type"], ["Optional columns: dataset_id (blank applies to all selected data), parameters (JSON object), severity, enabled, approved"], ["Use dataset_id from your selected datasets. Leave blank to make the rule job-wide."], ["Rule types: "+", ".join(sorted(RULE_TYPES))], ["Examples: RANGE parameters {\"min\":0,\"max\":100}; REGEX {\"pattern\":\"^[A-Z]+$\"}; IN_SET {\"values\":[\"A\",\"B\"]}"], ["Imported rules are validated and shown for review before they are saved."]]: info.append(row)
    info.column_dimensions["A"].width=115
    for c in info[1]: c.font=Font(bold=True,size=14,color="173B52")
    output=BytesIO(); wb.save(output); output.seek(0); return output


@app.get("/api/v1/rules/template")
def download_rules_template():
    return StreamingResponse(_rules_workbook(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition":"attachment; filename=data-quality-rules-template.xlsx"})


@app.post("/api/v1/jobs/{job_id}/rules/import-preview")
async def preview_rule_import(job_id: int, file: UploadFile = File(...), db: Session = Depends(get_db)):
    job = not_found(db.get(Job, job_id), "Job")
    name=(file.filename or "").lower()
    if not name.endswith((".xlsx", ".xls")): raise HTTPException(415, "Upload an .xlsx or .xls workbook")
    raw=await file.read()
    try:
        if name.endswith(".xlsx"):
            from openpyxl import load_workbook
            sheet=load_workbook(BytesIO(raw), read_only=True, data_only=True).active
            matrix=list(sheet.iter_rows(values_only=True))
        else:
            import xlrd
            book=xlrd.open_workbook(file_contents=raw); sheet=book.sheet_by_index(0)
            matrix=[[sheet.cell_value(r,c) for c in range(sheet.ncols)] for r in range(sheet.nrows)]
    except Exception as e: raise HTTPException(422, "Could not read the workbook. Check that it is a valid Excel file.") from e
    if not matrix: raise HTTPException(422, "The Rules sheet is empty")
    headers=[str(v or "").strip().lower() for v in matrix[0]]
    if "column_name" not in headers or "rule_type" not in headers: raise HTTPException(422, "Required columns column_name and rule_type are missing")
    existing=list(db.scalars(select(Rule).where(Rule.job_id==job_id)))
    known={rule_signature(r) for r in existing}; seen=set(); valid=[]; errors=[]
    datasets={d.id:d for d in db.scalars(select(Dataset).where(Dataset.id.in_(job.dataset_ids))).all()}
    for line, values in enumerate(matrix[1:], 2):
        row={headers[i]: values[i] if i<len(values) else None for i in range(len(headers))}
        if not any(v is not None and str(v).strip() for v in row.values()): continue
        try:
            column=str(row.get("column_name") or "").strip(); kind=str(row.get("rule_type") or "").strip().upper()
            if not column: raise ValueError("column_name is required")
            if not kind: raise ValueError("rule_type is required")
            dsval=row.get("dataset_id"); dsid=int(dsval) if dsval is not None and str(dsval).strip() else None
            if dsid is not None and dsid not in datasets: raise ValueError("dataset_id must belong to this job")
            if dsid is not None and column not in (datasets[dsid].schema_json or {}).get("columns",{}): raise ValueError("column_name is not present in the selected dataset")
            if dsid is None and not any(column in (d.schema_json or {}).get("columns",{}) for d in datasets.values()): raise ValueError("column_name was not found in selected datasets")
            params=row.get("parameters") or {}
            if isinstance(params,str): params=json.loads(params) if params.strip() else {}
            if not isinstance(params,dict): raise ValueError("parameters must be a JSON object")
            validate_rule_definition(kind,params)
            severity=str(row.get("severity") or "medium").strip().lower()
            if severity not in {"low","medium","high","critical"}: raise ValueError("severity must be low, medium, high, or critical")
            rule=Rule(job_id=job_id,dataset_id=dsid,column_name=column,rule_type=kind,parameters={**params,"severity":severity},enabled=str(row.get("enabled") if row.get("enabled") is not None else "false").lower() in ("true","1","yes"),approved=str(row.get("approved") if row.get("approved") is not None else "false").lower() in ("true","1","yes"))
            sig=rule_signature(rule)
            if sig in known or sig in seen: raise ValueError("Duplicate rule (already exists or repeated in this file)")
            seen.add(sig); valid.append({"row_number":line,"dataset_id":dsid,"column_name":column,"rule_type":kind,"parameters":rule.parameters,"enabled":rule.enabled,"approved":rule.approved,"severity":severity})
        except Exception as e: errors.append({"row_number":line,"error":str(e),"row":{k:(v.isoformat() if hasattr(v,"isoformat") else v) for k,v in row.items()}})
    return {"valid":valid,"errors":errors}


@app.post("/api/v1/jobs/{job_id}/rules/import")
def import_rules(job_id: int, body: dict, db: Session = Depends(get_db)):
    job=not_found(db.get(Job,job_id),"Job")
    saved=[]
    existing={rule_signature(r) for r in db.scalars(select(Rule).where(Rule.job_id==job_id))}
    for item in body.get("rules",[]):
        payload=RuleCreate(**item)
        kind=payload.rule_type.upper(); validate_rule_definition(kind,payload.parameters)
        if payload.dataset_id is not None:
            if payload.dataset_id not in job.dataset_ids: raise HTTPException(422,"Rule dataset_id must be selected by the job")
            ds=db.get(Dataset,payload.dataset_id)
            if not ds or payload.column_name not in (ds.schema_json or {}).get("columns",{}): raise HTTPException(422,"Rule column does not exist in the selected dataset")
        rule=Rule(job_id=job_id,**{**payload.model_dump(),"rule_type":kind})
        signature=rule_signature(rule)
        if signature in existing: raise HTTPException(409,"An imported rule is now a duplicate; preview the workbook again")
        existing.add(signature); db.add(rule); saved.append(rule)
    db.commit()
    for rule in saved: db.refresh(rule)
    return {"added":len(saved),"rules":saved}


@app.post("/api/v1/rules/error-report")
def download_rule_errors(body: dict):
    from openpyxl import Workbook
    wb=Workbook(); ws=wb.active; ws.title="Invalid rows"; ws.append(["Excel row","Error","Original row"])
    for item in body.get("errors",[]): ws.append([item.get("row_number"),item.get("error"),json.dumps(item.get("row",{}),default=str)])
    output=BytesIO(); wb.save(output); output.seek(0)
    return StreamingResponse(output,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":"attachment; filename=data-quality-rule-errors.xlsx"})


@app.post("/api/v1/jobs/{job_id}/baseline-rules", status_code=201)
def add_employee_baseline_rules(job_id: int, db: Session = Depends(get_db)):
    """Idempotently add the application's documented employee baseline checks.

    These fixed product rules are distinct from AI suggestions and are active
    only when the user explicitly adds them from the rule review screen.
    Existing rules and their approval decisions are never edited.
    """
    job = not_found(db.get(Job, job_id), "Job")
    if len(job.dataset_ids) != 1:
        raise HTTPException(422, "Employee baseline rules require a job with one dataset")
    dataset = not_found(db.get(Dataset, job.dataset_ids[0]), "Dataset")
    required = ["EmployeeID", "Name", "Email", "Age", "Department", "Salary"]
    schema_columns = (dataset.schema_json or {}).get("columns", {})
    if any(column not in schema_columns for column in required):
        raise HTTPException(422, "Dataset must contain EmployeeID, Name, Email, Age, Department, and Salary")
    selected = (job.selected_columns or {}).get(str(dataset.id), list(schema_columns))
    if any(column not in selected for column in required):
        raise HTTPException(422, "Select all six employee baseline columns before adding baseline rules")

    baseline_id = "employee-v1"
    definitions = [
        (column, "NOT_NULL", {}) for column in required
    ] + [
        ("Email", "EMAIL_FORMAT", {}),
        ("Age", "DATA_TYPE", {"type": "number"}),
        ("Age", "RANGE", {"min": 0}),
        ("Salary", "DATA_TYPE", {"type": "number"}),
        ("Salary", "RANGE", {"min": 0}),
        ("Department", "IN_SET", {"values": ["IT", "HR", "Finance", "Engineering", "Operations", "Marketing"], "case_sensitive": False}),
        ("EmployeeID", "UNIQUE", {}),
    ]
    existing = list(db.scalars(select(Rule).where(Rule.job_id == job.id)))
    known = {(r.dataset_id, r.column_name, r.rule_type.upper(),
              json.dumps({k:v for k,v in (r.parameters or {}).items() if k not in {"severity", "rationale"}}, sort_keys=True))
             for r in existing if (r.parameters or {}).get("baseline_id") == baseline_id}
    added = []
    for column, kind, params in definitions:
        marked = {**params, "baseline_id": baseline_id}
        key = (dataset.id, column, kind, json.dumps(marked, sort_keys=True))
        if key in known:
            continue
        rule = Rule(job_id=job.id, dataset_id=dataset.id, column_name=column, rule_type=kind,
                    parameters=marked, enabled=True, approved=True)
        db.add(rule)
        added.append(rule)
    db.commit()
    for rule in added:
        db.refresh(rule)
    return {"baseline_id": baseline_id, "added": len(added), "rules": added,
            "note": "Product-defined employee baseline rules are enabled and approved. AI suggestions still require review."}


@app.post("/api/v1/jobs/{job_id}/rules", status_code=201)
def create_rule(job_id: int, body: RuleCreate, db: Session = Depends(get_db)):
    job=not_found(db.get(Job, job_id), "Job")
    kind=body.rule_type.upper()
    try: validate_rule_definition(kind,body.parameters)
    except ValueError as e: raise HTTPException(422,str(e)) from e
    if body.dataset_id is not None:
        if body.dataset_id not in job.dataset_ids: raise HTTPException(422,"Rule dataset_id must be selected by the job")
        ds=db.get(Dataset,body.dataset_id)
        if not ds or body.column_name not in ds.schema_json.get("columns",{}): raise HTTPException(422,"Rule column does not exist in the selected dataset")
    rule = Rule(job_id=job_id, **{**body.model_dump(), "rule_type": kind})
    db.add(rule); db.commit(); db.refresh(rule); return rule


@app.patch("/api/v1/rules/{rule_id}")
def update_rule(rule_id: int, body: RuleUpdate, db: Session = Depends(get_db)):
    rule = not_found(db.get(Rule, rule_id), "Rule")
    changes = body.model_dump(exclude_unset=True)
    if any(k in changes and changes[k] is None for k in ("column_name","rule_type","parameters","enabled","approved")):
        raise HTTPException(422,"Rule fields cannot be null; omit fields you do not want to change")
    if "rule_type" in changes:
        changes["rule_type"] = changes["rule_type"].upper()
        if changes["rule_type"] not in RULE_TYPES: raise HTTPException(422, "Unsupported rule type")
    kind=changes.get("rule_type",rule.rule_type)
    params=changes.get("parameters",rule.parameters)
    try: validate_rule_definition(kind,params)
    except ValueError as e: raise HTTPException(422,str(e)) from e
    for k,v in changes.items(): setattr(rule,k,v)
    db.commit(); db.refresh(rule); return rule


@app.delete("/api/v1/rules/{rule_id}", status_code=204)
def delete_rule(rule_id: int, db: Session = Depends(get_db)):
    rule = not_found(db.get(Rule, rule_id), "Rule"); db.delete(rule); db.commit()


@app.post("/api/v1/jobs/{job_id}/suggest-rules")
async def suggest_rules(job_id: int, db: Session = Depends(get_db)):
    job = not_found(db.get(Job, job_id), "Job")
    datasets = [not_found(db.get(Dataset, did), "Dataset") for did in job.dataset_ids]
    suggestions = await generate_rule_suggestions(datasets, db, job.selected_columns or {})
    existing = list(db.scalars(select(Rule).where(Rule.job_id == job_id)))
    existing_keys = {rule_key(r.dataset_id, r.column_name, r.rule_type, r.parameters) for r in existing}
    added = []
    for item in suggestions:
        params = item["parameters"]
        key = rule_key(item["dataset_id"], item["column_name"], item["rule_type"], params)
        if key in existing_keys: continue
        existing_keys.add(key)
        rule = Rule(job_id=job.id, dataset_id=item["dataset_id"], column_name=item["column_name"],
                    rule_type=item["rule_type"], parameters={**params, "severity": item["severity"], "rationale": item["rationale"]},
                    enabled=True, approved=False)
        db.add(rule); added.append(rule)
    db.commit()
    for rule in added: db.refresh(rule)
    return {"provider":settings.ai_provider,"model":settings.resolved_ai_model,"suggestions":added,"saved":True,"added":len(added),"note":"AI-generated rules are pending review and approval."}


@app.post("/api/v1/datasets/{dataset_id}/suggest-rules")
async def suggest_dataset_rules(dataset_id: int, db: Session = Depends(get_db)):
    dataset = not_found(db.get(Dataset, dataset_id), "Dataset")
    suggestions = await generate_rule_suggestions([dataset], db)
    return {"provider":settings.ai_provider,"model":settings.resolved_ai_model,"suggestions":suggestions,"saved":False}


def rule_key(dataset_id, column_name, rule_type, parameters):
    relevant = {k:v for k,v in (parameters or {}).items() if k not in {"severity", "rationale"}}
    return (dataset_id, column_name, rule_type.upper(), json.dumps(relevant, sort_keys=True, default=str))


def _negative_age_range_suggestions(profile, suggestions):
    """Add a deterministic non-negative bound when observed age data is negative."""
    age_names = {"age", "customer_age", "age_years", "years_old"}
    for dataset in profile:
        for column, metadata in dataset.get("columns", {}).items():
            if column.lower() not in age_names:
                continue
            if str(metadata.get("inferred_type", "")).lower() not in {"int", "integer", "float", "number", "decimal"}:
                continue
            try:
                has_negative = any(float(value) < 0 for value in metadata.get("distinct_sample", []))
            except (TypeError, ValueError):
                continue
            if not has_negative:
                continue
            has_nonnegative_bound = any(
                item.get("dataset_id") == dataset["dataset_id"]
                and item.get("column_name") == column
                and item.get("rule_type") == "RANGE"
                and item.get("parameters", {}).get("min") is not None
                and float(item["parameters"]["min"]) >= 0
                for item in suggestions
            )
            if not has_nonnegative_bound:
                suggestions.append({"dataset_id": dataset["dataset_id"], "column_name": column,
                    "rule_type": "RANGE", "parameters": {"min": 0}, "severity": "high",
                    "rationale": "Observed negative age values; age must be non-negative."})
    return suggestions


def _safe_provider_error(response, api_key):
    """Extract a short provider diagnostic without returning credential material."""
    try:
        body = response.json()
    except (ValueError, TypeError):
        body = None
    error = body.get("error") if isinstance(body, dict) else None
    if isinstance(error, dict):
        detail = error.get("message") or error.get("type") or error.get("code")
    elif isinstance(error, str):
        detail = error
    else:
        detail = None
    if not detail:
        detail = response.text
    detail = str(detail)
    detail = re.sub(r"(?i)bearer\s+\S+", "Bearer [redacted]", detail)
    detail = re.sub(r"(?i)(api[_ -]?key|access[_ -]?token|client[_ -]?secret)\s*[:=]\s*\S+", r"\1=[redacted]", detail)
    detail = re.sub(r"(?i)\b(gsk_[A-Za-z0-9_-]+|sk-[A-Za-z0-9_-]+)\b", "[redacted]", detail)
    if api_key:
        detail = detail.replace(api_key, "[redacted]")
    return detail[:500].strip() or "No error details returned"


async def generate_rule_suggestions(datasets, db, selected_columns=None):
    ai_base_url=settings.resolved_ai_base_url
    ai_api_key=settings.resolved_ai_api_key
    ai_model=settings.resolved_ai_model
    if settings.ai_provider == "none" or not ai_api_key or not ai_model or not ai_base_url:
        raise HTTPException(503, "AI rule suggestions are disabled; configure the AI provider, API key, model, and endpoint")
    import httpx
    profile = []
    local = LocalStorage()
    for ds in datasets:
        schema = ds.schema_json or {}
        allowed = (selected_columns or {}).get(str(ds.id))
        cols = {c:dict(meta) for c,meta in schema.get("columns",{}).items() if allowed is None or c in allowed}
        # Profile local file rows where available. Connector datasets already carry schema metadata;
        # never assume their remote paths can be read through local storage.
        if ds.source_type in {"local", "upload"}:
            try:
                rows = parse_data(ds.filename, local.read(ds.path))
                for c, meta in cols.items():
                    vals = [r.get(c) for r in rows]
                    meta.update({"inferred_type":infer_type(vals),"null_fraction":sum(v in (None, "") for v in vals)/max(len(vals),1),
                                 "distinct_sample":list(dict.fromkeys(str(v) for v in vals if v not in (None,"")))[:10]})
                schema["row_count"] = len(rows)
            except (OSError, ValueError):
                pass
        profile.append({"dataset_id": ds.id, "dataset_name": ds.name, "source_type": ds.source_type,
                        "row_count": schema.get("row_count"), "metadata":{k:v for k,v in schema.items() if k != "columns"}, "columns": cols})
    prompt = {"task":"Suggest relevant data quality rules grounded only in the provided schema metadata and profile. Return JSON with a suggestions array. Each item: dataset_id, column_name, rule_type, parameters, severity (low, medium, high, critical), rationale. Use only real columns and supported rule types. Inspect distinct_sample values for observed anomalies: when a semantically non-negative numeric field (for example age) contains a negative sample, include a RANGE rule with min 0; do not treat DATA_TYPE alone as sufficient. Do not guess other ranges or accepted values unless supported by profile evidence. Use NOT_NULL, UNIQUE, DATA_TYPE, RANGE, EMAIL_FORMAT, REGEX, LENGTH, or IN_SET as appropriate.", "supported_rule_types":sorted(RULE_TYPES), "profile":profile}
    url = ai_base_url.rstrip("/") + "/chat/completions"
    try:
        async with httpx.AsyncClient(timeout=settings.ai_timeout_seconds) as client:
            response = await client.post(url, headers={"Authorization":f"Bearer {ai_api_key}"}, json={"model":ai_model,"messages":[{"role":"user","content":json.dumps(prompt)}],"temperature":0.2,"response_format":{"type":"json_object"}})
        response.raise_for_status(); payload=response.json()
        text=payload["choices"][0]["message"]["content"]
        result=json.loads(text)
        if isinstance(result,list): suggestions=result
        elif isinstance(result,dict): suggestions=result.get("suggestions",[])
        else: raise ValueError("Expected an object with a suggestions array")
        if not isinstance(suggestions,list): raise ValueError("AI response suggestions field must be an array")
        known = {ds.id:set((ds.schema_json or {}).get("columns",{})) & set((selected_columns or {}).get(str(ds.id), (ds.schema_json or {}).get("columns",{}))) for ds in datasets}
        clean=[]
        if len(suggestions) > 200: raise ValueError("Too many suggestions")
        for item in suggestions:
            if not isinstance(item,dict): raise ValueError("Suggestion must be an object")
            did=int(item.get("dataset_id", -1)); column=item.get("column_name"); kind=str(item.get("rule_type", "")).upper()
            params=item.get("parameters", {})
            if did not in known or column not in known[did] or kind not in RULE_TYPES or not isinstance(params,dict):
                raise ValueError("AI response referenced an unknown column or unsupported rule")
            validate_rule_definition(kind, params)
            severity=str(item.get("severity", "medium")).lower()
            if severity not in {"low","medium","high","critical"}: raise ValueError("Invalid severity")
            clean.append({"dataset_id":did,"column_name":column,"rule_type":kind,"parameters":params,
                          "severity":severity,"rationale":str(item.get("rationale", "Suggested from schema metadata"))[:500]})
        return _negative_age_range_suggestions(profile, clean)
    except httpx.HTTPStatusError as e:
        status = e.response.status_code
        request_id = e.response.headers.get("x-request-id") or e.response.headers.get("x-groq-id")
        provider_detail = _safe_provider_error(e.response, ai_api_key)
        logging.getLogger(__name__).warning(
            "AI provider returned HTTP %s (provider=%s model=%s request_id=%s detail=%s)",
            status, settings.ai_provider, ai_model, request_id or "unavailable", provider_detail,
        )
        if status == 401:
            raise HTTPException(502, f"AI provider returned HTTP 401 (invalid API key): {provider_detail}") from e
        if status == 403:
            raise HTTPException(502, f"AI provider returned HTTP 403 (permission or model access issue): {provider_detail}") from e
        if status == 429:
            raise HTTPException(503, f"AI provider returned HTTP 429 (rate limit or quota issue): {provider_detail}") from e
        if status == 400:
            raise HTTPException(502, f"AI provider returned HTTP 400 (invalid request or unsupported parameter): {provider_detail}") from e
        if status == 404:
            raise HTTPException(502, f"AI provider returned HTTP 404 (endpoint or model unavailable): {provider_detail}") from e
        if status == 422:
            raise HTTPException(502, f"AI provider returned HTTP 422 (request validation failed): {provider_detail}") from e
        if status >= 500:
            raise HTTPException(503, f"AI provider returned HTTP {status} (provider-side failure): {provider_detail}") from e
        raise HTTPException(502, f"AI provider returned HTTP {status}: {provider_detail}") from e
    except httpx.TimeoutException as e:
        logging.getLogger(__name__).warning("AI provider request timed out (provider=%s model=%s)", settings.ai_provider, ai_model)
        raise HTTPException(504, "AI provider request timed out; try again") from e
    except httpx.RequestError as e:
        logging.getLogger(__name__).warning(
            "AI provider connection failed (provider=%s model=%s error=%s)",
            settings.ai_provider, ai_model, type(e).__name__,
        )
        raise HTTPException(502, "Could not connect to the AI provider; verify the endpoint and network access") from e
    except (KeyError, ValueError, TypeError) as e:
        logging.getLogger(__name__).warning(
            "AI provider returned an invalid response (provider=%s model=%s error=%s)",
            settings.ai_provider, ai_model, type(e).__name__,
        )
        raise HTTPException(502, "AI provider returned an invalid response; no rules were created") from e


@app.post("/api/v1/jobs/{job_id}/execute", status_code=202)
def run_job(job_id: int, tasks: BackgroundTasks, db: Session = Depends(get_db)):
    not_found(db.get(Job, job_id), "Job")
    active = db.scalar(select(Execution.id).where(
        Execution.job_id == job_id, Execution.status.in_(("queued", "running"))
    ).limit(1))
    if active is not None:
        raise HTTPException(409, f"Job already has an active execution ({active})")
    record = Execution(job_id=job_id, status="queued", rule_snapshot=rule_snapshot(db, job_id))
    db.add(record); db.commit(); db.refresh(record)
    tasks.add_task(execute, record.id)
    return record


@app.post("/api/v1/executions/{execution_id}/retry", status_code=202)
def retry_execution(execution_id: int, tasks: BackgroundTasks, db: Session = Depends(get_db)):
    previous=not_found(db.get(Execution,execution_id),"Execution")
    if previous.status != "failed": raise HTTPException(409,"Only failed executions can be retried")
    active = db.scalar(select(Execution.id).where(
        Execution.job_id == previous.job_id, Execution.status.in_(("queued", "running"))
    ).limit(1))
    if active is not None:
        raise HTTPException(409, f"Job already has an active execution ({active})")
    retry=Execution(job_id=previous.job_id,status="queued",attempt=previous.attempt+1,retry_of=previous.id,
                    rule_snapshot=previous.rule_snapshot if previous.rule_snapshot is not None else rule_snapshot(db, previous.job_id))
    db.add(retry);db.commit();db.refresh(retry);tasks.add_task(execute,retry.id);return retry


@app.get("/api/v1/jobs/{job_id}/executions")
def job_executions(job_id: int, db: Session = Depends(get_db)):
    not_found(db.get(Job, job_id), "Job")
    return db.scalars(select(Execution).where(Execution.job_id == job_id).order_by(Execution.id.desc())).all()


def rule_snapshot(db: Session, job_id: int):
    """Capture every rule and its lifecycle state when a run is requested."""
    job = not_found(db.get(Job, job_id), "Job")
    scope = [Rule.dataset_id.is_(None)]
    if job.dataset_ids:
        scope.append(Rule.dataset_id.in_(job.dataset_ids))
    rules = db.scalars(select(Rule).where(Rule.job_id == job_id, or_(*scope)).order_by(Rule.id)).all()
    return [{"id": r.id, "job_id": r.job_id, "dataset_id": r.dataset_id,
             "column_name": r.column_name, "rule_type": r.rule_type,
             "parameters": r.parameters or {}, "enabled": bool(r.enabled),
             "approved": bool(r.approved)} for r in rules]


@app.get("/api/v1/executions/{execution_id}")
def execution_detail(execution_id: int, tasks: BackgroundTasks, db: Session = Depends(get_db)):
    execution = not_found(db.get(Execution, execution_id), "Execution")
    # BackgroundTasks are process local. A queued row may outlive its request
    # task (for example, after a worker restart). Polling the persisted row
    # safely re-dispatches it; execute() atomically claims queued rows, so
    # duplicate dispatches cannot run the same execution twice.
    if execution.status == "queued":
        tasks.add_task(execute, execution.id)
    return execution


@app.get("/api/v1/executions/{execution_id}/failed")
def failed_records(execution_id: int, db: Session = Depends(get_db)):
    e=not_found(db.get(Execution, execution_id), "Execution")
    if not e.report_path: raise HTTPException(409, "Execution has no saved outputs")
    job=not_found(db.get(Job,e.job_id),"Job")
    destination=destination_for_job(job)
    if destination.get("type")=="databricks":
        report=json.loads(LocalStorage().read(e.report_path))
        result=report.get("destination_result") or {}
        if result.get("type")!="databricks" or not result.get("failed_table"):
            raise HTTPException(409,"Execution has no saved Databricks FAIL table")
        secrets={}
        credential=db.get(Credential,destination.get("credential_id")) if destination.get("credential_id") else None
        if credential:
            if not settings.secret_encryption_key: raise HTTPException(500,"SECRET_ENCRYPTION_KEY is required for encrypted cloud credentials")
            secrets=json.loads(Fernet(settings.secret_encryption_key.encode()).decrypt(credential.secret_ciphertext.encode()))
            destination={**(credential.config or {}),**destination}
        else:
            destination={**destination,"server_hostname":destination.get("server_hostname") or settings.databricks_server_hostname,
                "http_path":destination.get("http_path") or settings.databricks_http_path,
                "catalog":result.get("catalog") or destination.get("catalog") or settings.databricks_catalog,
                "schema":result.get("schema") or destination.get("schema") or settings.databricks_schema}
            secrets={"access_token":settings.databricks_token} if settings.databricks_token else {}
        try:
            remote=read_databricks_table(destination,secrets,result["failed_table"])
        except Exception as exc:
            raise HTTPException(502,"Could not load failed records from the Databricks execution table") from exc
        output=[]
        for row in remote:
            index=row.get("_dq_row_index")
            if index is None: raise HTTPException(502,"Databricks FAIL table is missing original row indexes")
            reasons=row.get("_dq_failure_reasons")
            try: reasons=json.loads(reasons) if isinstance(reasons,str) else reasons
            except (TypeError,ValueError): reasons=[]
            record={k:v for k,v in row.items() if not k.startswith("_dq_")}
            output.append({"row_index":int(index),"record":record,"failure_reasons":reasons or []})
        by_index={item["row_index"]:item for item in output}
        for correction in db.scalars(select(Correction).where(Correction.execution_id==execution_id).order_by(Correction.id)):
            if correction.row_index in by_index:
                by_index[correction.row_index]["record"][correction.column_name]=correction.new_value
        return output
    path=Path(e.report_path).parent / "failed.json"
    try: return json.loads(LocalStorage().read(str(path)))
    except FileNotFoundError: raise HTTPException(404, "Failed records output not found")


@app.get("/api/v1/executions/{execution_id}/artifacts/{artifact_name}")
def execution_artifact(execution_id: int, artifact_name: str, db: Session = Depends(get_db)):
    e=not_found(db.get(Execution, execution_id), "Execution")
    if not e.report_path: raise HTTPException(409, "Execution has no saved outputs")
    if artifact_name not in {"passed.csv", "failed.json", "report.json", "revalidated_passed.csv", "revalidated_failed.json"}:
        raise HTTPException(404,"Artifact not found")
    path=Path(e.report_path).parent / artifact_name
    try: content=LocalStorage().read(str(path))
    except FileNotFoundError: raise HTTPException(404,"Artifact not found")
    media="text/csv" if artifact_name.endswith(".csv") else "application/json"
    return JSONResponse(content=json.loads(content)) if artifact_name.endswith(".json") else __import__("fastapi").Response(content=content,media_type=media,headers={"Content-Disposition":f'attachment; filename="{artifact_name}"'})


@app.post("/api/v1/executions/{execution_id}/corrections")
def correct_record(execution_id: int, body: CorrectionCreate, db: Session = Depends(get_db)):
    e=not_found(db.get(Execution, execution_id), "Execution")
    if not e.report_path: raise HTTPException(409, "Execution has no saved outputs")
    path=Path(e.report_path).parent / "failed.json"
    rows=json.loads(LocalStorage().read(str(path)))
    item=next((r for r in rows if r["row_index"] == body.row_index), None)
    if item is None: raise HTTPException(404, "Failed row not found")
    if body.column_name not in item["record"]: raise HTTPException(422, "Column does not exist in failed record")
    old=item["record"][body.column_name]
    corr=Correction(execution_id=execution_id,row_index=body.row_index,column_name=body.column_name,old_value=None if old is None else str(old),new_value=body.value)
    item["record"][body.column_name]=body.value
    LocalStorage().save(str(path.relative_to(Path(settings.local_data_dir).resolve())), json.dumps(rows,ensure_ascii=False,indent=2).encode())
    db.add(corr); db.commit()
    return {"correction_id":corr.id,"row_index":body.row_index,"column_name":body.column_name,"old_value":old,"new_value":body.value}


@app.post("/api/v1/executions/{execution_id}/revalidate")
def revalidate(execution_id: int, db: Session = Depends(get_db)):
    e=not_found(db.get(Execution, execution_id), "Execution")
    if not e.report_path: raise HTTPException(409, "Execution has no saved outputs")
    path=Path(e.report_path).parent / "failed.json"
    rows=json.loads(LocalStorage().read(str(path)))
    job=db.get(Job,e.job_id)
    destination=destination_for_job(job)
    # failed.json is the execution-scoped correction source of truth. The
    # failed-record endpoint may read the Databricks table for display, but
    # revalidation must not read it again: that extra SQL session duplicated
    # work immediately before the write and discarded already saved edits.
    # Overlay history so corrections saved by older clients are honored too.
    by_index={item["row_index"]:item for item in rows}
    for correction in db.scalars(select(Correction).where(Correction.execution_id==execution_id).order_by(Correction.id)):
        if correction.row_index in by_index:
            by_index[correction.row_index]["record"][correction.column_name]=correction.new_value
    rows=list(by_index.values())
    records=[r["record"] for r in rows]
    if e.rule_snapshot is not None:
        from types import SimpleNamespace
        rules=[SimpleNamespace(**item) for item in e.rule_snapshot]
    else:
        scope = [Rule.dataset_id.is_(None)]
        if job.dataset_ids:
            scope.append(Rule.dataset_id.in_(job.dataset_ids))
        rules=list(db.scalars(select(Rule).where(Rule.job_id==job.id, or_(*scope))))
    selected_columns=job.selected_columns or {}
    def applies(rule):
        if not rule.enabled or not rule.approved: return False
        if rule.dataset_id is not None:
            columns=selected_columns.get(str(rule.dataset_id), selected_columns.get(rule.dataset_id, []))
            return not columns or rule.column_name in columns
        return not selected_columns or any(rule.column_name in columns for columns in selected_columns.values())
    rules=[rule for rule in rules if applies(rule)]
    unique_rules=[]; seen_signatures=set()
    for rule in rules:
        signature=rule_signature(rule)
        if signature not in seen_signatures: unique_rules.append(rule); seen_signatures.add(signature)
    rules=unique_rules
    from app.rules import validate_records
    passed,failed=validate_records(records,rules)
    # validate_records indexes the supplied subset; restore the original source
    # row positions so retries/revalidation replace the same execution rows.
    newly_failed_positions={item["row_index"] for item in failed}
    newly_passed_indices=[rows[index]["row_index"] for index in range(len(rows))
                          if index not in newly_failed_positions]
    for item in failed:
        item["row_index"] = rows[item["row_index"]]["row_index"]
    root=Path(settings.local_data_dir).resolve(); base=Path(e.report_path).parent
    local=LocalStorage()
    report=json.loads(local.read(e.report_path))
    prior_failed_indices={item["row_index"] for item in rows}
    passed_indices=[index for index in range(report.get("records_loaded",e.passed_count+e.failed_count))
                    if index not in prior_failed_indices]
    original_passed=parse_data("passed.csv", local.read(str(base/"passed.csv")))
    all_passed=original_passed+passed
    all_failed=failed
    destination=destination_for_job(job)
    provider=destination.get("type","local")
    secrets={}
    if provider != "local":
        credential_id=destination.get("credential_id")
        credential=db.get(Credential,credential_id) if credential_id else None
        if credential:
            if not settings.secret_encryption_key:
                raise HTTPException(500,"SECRET_ENCRYPTION_KEY is required for encrypted cloud credentials")
            secrets=json.loads(Fernet(settings.secret_encryption_key.encode()).decrypt(credential.secret_ciphertext.encode()))
            destination={**(credential.config or {}),**destination}
        elif provider == "databricks":
            destination={**destination,
                "server_hostname":destination.get("server_hostname") or settings.databricks_server_hostname,
                "http_path":destination.get("http_path") or settings.databricks_http_path,
                "catalog":destination.get("catalog") or settings.databricks_catalog,
                "schema":destination.get("schema") or settings.databricks_schema}
            secrets={"access_token":settings.databricks_token} if settings.databricks_token else {}
        else:
            raise HTTPException(422,f"Destination {provider} requires a saved credential")
    source_datasets=[db.get(Dataset,dataset_id) for dataset_id in (job.dataset_ids or [])]
    dataset_label="_and_".join(item.name for item in source_datasets if item)
    storage_run_id=e.id
    parent=db.get(Execution,e.retry_of) if e.retry_of else None
    while parent:
        storage_run_id=parent.id
        parent=db.get(Execution,parent.retry_of) if parent.retry_of else None
    try:
        remote=write_destination(destination,secrets,str(base),all_passed,all_failed,report,encode_csv(all_passed),
            dataset_name=dataset_label or job.name,job_id=job.id,execution_id=e.id,
            columns=list(dict.fromkeys(k for row in all_passed for k in row)),
            source_dataset_ids=[item.id for item in source_datasets if item],storage_run_id=storage_run_id,
            passed_row_indices=passed_indices + newly_passed_indices,
            atomic_update=provider=="databricks")
        report["destination_result"]=remote
        report["storage_status"]={"status":"saved","pass_status":remote.get("pass_status","saved"),"fail_status":remote.get("fail_status","saved")}
        e.status="succeeded";e.error=None
    except Exception as exc:
        # Do not publish changed local outcomes or counts when the destination
        # could not persist them. The failed artifact remains the retry source.
        db.rollback()
        raise HTTPException(502, "Destination storage failed; corrections remain saved. Retry revalidation after fixing the destination.") from exc
    local.save(str((base/"revalidated_passed.csv").relative_to(root)),encode_csv(passed))
    local.save(str((base/"revalidated_failed.json").relative_to(root)),json.dumps(failed,ensure_ascii=False,indent=2).encode())
    passed_path=local.save(str((base/"passed.csv").relative_to(root)),encode_csv(all_passed))
    failed_path=local.save(str((base/"failed.json").relative_to(root)),json.dumps(all_failed,ensure_ascii=False,indent=2).encode())
    report.update(status=e.status,passed=len(all_passed),failed=len(all_failed),
                   total=len(all_passed)+len(all_failed),records_evaluated=len(all_passed)+len(all_failed),
                   passed_path=passed_path,failed_path=failed_path)
    output_dataset=db.get(Dataset,report.get("output_dataset_id"))
    if output_dataset:
        output_dataset.schema_json={"columns":{column:{"inferred_type":"string","nullable":any(row.get(column) in (None,"") for row in all_passed)} for column in dict.fromkeys(k for row in all_passed for k in row)},"row_count":len(all_passed)}
    local.save(str((base/"report.json").relative_to(root)),json.dumps(report,ensure_ascii=False,indent=2).encode())
    e.passed_count=len(all_passed);e.failed_count=len(all_failed)
    if not all_failed and e.status=="succeeded": e.error=None
    e.logs=(e.logs or [])+[f"Revalidated {len(records)} corrected records; {len(all_passed)} passed; {len(all_failed)} failed",f"Destination storage {report['storage_status']['status']}"]
    db.commit()
    return {"revalidated":len(records),"passed":len(passed),"failed":len(failed),
            "total_passed":len(all_passed),"total_failed":len(all_failed),"storage_status":report["storage_status"]}


@app.get("/api/v1/executions/{execution_id}/corrections")
def correction_history(execution_id: int, db: Session = Depends(get_db)):
    not_found(db.get(Execution,execution_id),"Execution")
    return db.scalars(select(Correction).where(Correction.execution_id==execution_id).order_by(Correction.id)).all()


@app.post("/api/v1/credentials", status_code=201)
def create_credential(body: CredentialCreate, db: Session = Depends(get_db)):
    sensitive=sensitive_config_keys(body.config)
    if sensitive: raise HTTPException(422,f"Sensitive values belong in secrets, not config: {sensitive}")
    cipher=credential_key()
    item=Credential(name=body.name,provider=body.provider,secret_ciphertext=cipher.encrypt(json.dumps(body.secrets).encode()).decode(),config=body.config)
    db.add(item); db.commit(); db.refresh(item)
    return {"id":item.id,"name":item.name,"provider":item.provider,"config":item.config,"created_at":item.created_at}


@app.post("/api/v1/sources/connect", status_code=201)
def connect_source(body: CredentialCreate, prefix: str = "", db: Session = Depends(get_db)):
    """Validate a cloud source by listing its input prefix before storing credentials."""
    provider = "onelake" if body.provider == "fabric" else body.provider
    if provider not in ("s3", "azure_blob", "onelake", "databricks", "snowflake"):
        raise HTTPException(400, "Choose a supported data source")
    if provider == "s3":
        auth_method = body.config.get("auth_method", "access_key")
        if auth_method not in ("access_key", "default_chain"):
            raise HTTPException(422, "Choose access key or IAM role authentication")
        if not body.config.get("bucket") or not body.config.get("region"):
            raise HTTPException(422, "Enter the S3 bucket name and AWS region")
        if auth_method == "access_key" and (not body.secrets.get("access_key_id") or not body.secrets.get("secret_access_key")):
            raise HTTPException(422, "Enter an access key ID and secret access key")
        if auth_method == "default_chain" and (body.secrets.get("access_key_id") or body.secrets.get("secret_access_key")):
            raise HTTPException(422, "Remove access keys when using IAM role authentication")
    adapter = None
    try:
        adapter = get_source_adapter(provider, body.config, body.secrets)
        files = adapter.list_files(prefix)
        if files and provider != "s3":
            adapter.check_read(files[0]["name"])
    except Exception as e:
        if provider == "s3":
            response = getattr(e, "response", {})
            code = str(response.get("Error", {}).get("Code", "")) if isinstance(response, dict) else ""
            if code in ("InvalidAccessKeyId", "SignatureDoesNotMatch", "InvalidToken", "ExpiredToken", "UnrecognizedClientException"):
                raise HTTPException(401, "AWS rejected the supplied credentials. Check the access key ID, secret key, and session token.") from e
            if code in ("AccessDenied", "403", "Forbidden"):
                raise HTTPException(403, "AWS credentials do not have permission to inspect this bucket. Grant s3:ListBucket and bucket access.") from e
            if code in ("NoSuchBucket", "404", "NotFound"):
                raise HTTPException(404, "The S3 bucket was not found. Check the bucket name and AWS region.") from e
            if code in ("PermanentRedirect", "301", "AuthorizationHeaderMalformed"):
                raise HTTPException(400, "The bucket is in a different AWS region. Select its correct region.") from e
            if isinstance(e, RuntimeError) and "boto3" in str(e).lower():
                raise HTTPException(503, "S3 support is unavailable. Install the cloud dependencies (pip install .[cloud]) and restart the backend.") from e
            raise HTTPException(400, "Could not connect to S3; verify the bucket, region, and AWS credentials.") from e
        raise HTTPException(400, "Could not connect or list files; verify the source settings and credentials") from e
    finally:
        if adapter and hasattr(adapter, "close"): adapter.close()
    sensitive = sensitive_config_keys(body.config)
    if sensitive:
        raise HTTPException(422, f"Sensitive values belong in secrets, not config: {sensitive}")
    cipher = credential_key()
    for existing in db.scalars(select(Credential).where(Credential.provider == body.provider, Credential.name == body.name)).all():
        try:
            existing_secrets = json.loads(cipher.decrypt(existing.secret_ciphertext.encode()))
        except (InvalidToken, ValueError):
            continue
        if existing.config == body.config and existing_secrets == body.secrets:
            return {"id": existing.id, "name": existing.name, "provider": existing.provider, "files": files}
    item = Credential(name=body.name, provider=body.provider, secret_ciphertext=cipher.encrypt(json.dumps(body.secrets).encode()).decode(), config=body.config)
    db.add(item); db.commit(); db.refresh(item)
    return {"id": item.id, "name": item.name, "provider": item.provider, "files": files}


@app.get("/api/v1/credentials")
def list_credentials(db: Session = Depends(get_db)):
    return [{"id":x.id,"name":x.name,"provider":x.provider,"config":x.config,"created_at":x.created_at} for x in db.scalars(select(Credential)).all()]


@app.post("/api/v1/datasets/import/{credential_id}", status_code=201)
def import_cloud_dataset(credential_id: int, object_key: str, name: str | None = None, db: Session = Depends(get_db)):
    item=not_found(db.get(Credential,credential_id),"Credential")
    for existing in db.scalars(select(Dataset).where(Dataset.source_type == item.provider)).all():
        source_meta = (existing.schema_json or {}).get("_source", {})
        if source_meta.get("credential_id") == credential_id and source_meta.get("object_key") == object_key:
            return existing
    try: secrets=json.loads(credential_key().decrypt(item.secret_ciphertext.encode()))
    except (InvalidToken, ValueError) as e: raise HTTPException(500,"Stored credential cannot be decrypted") from e
    provider=item.provider
    try:
        storage_provider="onelake" if provider == "fabric" else provider
        if storage_provider not in ("s3","azure_blob","onelake","databricks","snowflake"): raise HTTPException(501,f"{provider} source adapter is not implemented")
        adapter=get_source_adapter(storage_provider,item.config,secrets)
        try: content=adapter.read(object_key)
        finally:
            if hasattr(adapter,"close"): adapter.close()
        if storage_provider in ("databricks","snowflake") and not object_key.lower().endswith((".csv",".tsv",".json",".xlsx",".xls",".parquet")):
            fname=Path(object_key).name + ".csv"
        else: fname=Path(object_key).name
        records=parse_data(fname if storage_provider in ("databricks","snowflake") else object_key,content)
    except Exception as e: raise HTTPException(502,"Source read failed; verify the connection and selected object") from e
    columns=list(dict.fromkeys(k for r in records[:1000] for k in r))
    schema={"columns":{c:{"inferred_type":infer_type([r.get(c) for r in records[:1000]]),"nullable":any(r.get(c) in (None,"") for r in records[:1000])} for c in columns},"row_count":len(records),"_source":{"credential_id":credential_id,"object_key":object_key}}
    path=LocalStorage().save(f"uploads/{__import__('uuid').uuid4().hex}-{fname}",content)
    ds=Dataset(name=name or Path(fname).stem,filename=fname,path=path,source_type=provider,schema_json=schema)
    db.add(ds);db.commit();db.refresh(ds);return ds


@app.get("/api/v1")
def api_root(): return {"name":"ADQ Automation API","docs":"/docs","version":"0.1.0"}
