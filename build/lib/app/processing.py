import csv, io, json
from datetime import datetime, timezone
from pathlib import Path
import logging
from types import SimpleNamespace
from sqlalchemy import or_, select
from app.config import settings
from app.connectors.storage import LocalStorage
from app.connectors.destinations import write_destination
from app.database import SessionLocal
from app.models import Credential, Dataset, Execution, Job, Rule
from app.rules import is_missing, rule_signature, validate_records


def parse_data(filename: str, content: bytes):
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        text = content.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text)))
    if suffix == ".tsv":
        text = content.decode("utf-8-sig")
        return list(csv.DictReader(io.StringIO(text), delimiter="\t"))
    if suffix == ".json":
        value = json.loads(content)
        if isinstance(value, dict): value = value.get("records", [value])
        if not isinstance(value, list) or any(not isinstance(x, dict) for x in value):
            raise ValueError("JSON must be an array of objects or an object with a records array")
        return value
    if suffix in (".xlsx", ".xls"):
        try:
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except ImportError as e:
            raise ValueError("Excel support requires openpyxl") from e
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        raw_headers = list(next(rows, []))
        data_rows = list(rows)
        headers = []
        keep = []
        for index, value in enumerate(raw_headers):
            # Empty Excel columns are common in exported workbooks. Do not
            # create synthetic fields for them unless the column has data.
            populated = any(index < len(row) and row[index] is not None for row in data_rows)
            if value is None or not str(value).strip():
                if not populated:
                    continue
                header = f"column_{index + 1}"
            else:
                header = str(value).strip()
            headers.append(header)
            keep.append(index)
        workbook.close()
        return [{headers[out_index]: row[index] for out_index, index in enumerate(keep) if index < len(row)}
                for row in data_rows if any(index < len(row) and row[index] is not None for index in keep)]
    if suffix == ".parquet":
        try:
            import pyarrow.parquet as parquet
            return parquet.read_table(io.BytesIO(content)).to_pylist()
        except ImportError as e:
            raise ValueError("Parquet support requires pyarrow") from e
    raise ValueError("Supported formats are CSV, TSV, JSON, Excel, and Parquet")


def encode_csv(records: list[dict]) -> bytes:
    if not records: return b""
    cols = list(dict.fromkeys(k for row in records for k in row))
    out = io.StringIO(newline="")
    writer = csv.DictWriter(out, fieldnames=cols, extrasaction="ignore")
    writer.writeheader(); writer.writerows(records)
    return out.getvalue().encode("utf-8")


def schema_from_records(records: list[dict]) -> dict:
    """Build schema metadata from the rows actually read for this execution."""
    columns = list(dict.fromkeys(key for row in records for key in row))
    metadata = {}
    for column in columns:
        values = [row.get(column) for row in records]
        present = [str(value).strip() for value in values if not is_missing(value)]
        if present and all(_is_integer(value) for value in present):
            kind = "integer"
        elif present and all(_is_number(value) for value in present):
            kind = "number"
        elif present and all("@" in value for value in present):
            kind = "email-like"
        else:
            kind = "string"
        metadata[column] = {"inferred_type": kind,
                            "nullable": any(is_missing(value) for value in values)}
    return {"columns": metadata, "row_count": len(records)}


def _is_integer(value):
    try:
        int(value)
        return True
    except (ValueError, TypeError):
        return False


def _is_number(value):
    try:
        float(value)
        return True
    except (ValueError, TypeError):
        return False


def validate_rule_bindings(rules, job_dataset_ids, records_by_dataset, combined_records):
    """Reject stale dataset bindings and columns absent from actual input rows."""
    from app.rules import resolve_column
    allowed = set(job_dataset_ids or [])
    for rule in rules:
        if rule.dataset_id is not None and rule.dataset_id not in allowed:
            raise ValueError(f"Rule {rule.id} references dataset {rule.dataset_id}, which is not selected by this job")
        rows = records_by_dataset.get(rule.dataset_id, []) if rule.dataset_id is not None else combined_records
        if rows:
            resolve_column(rows[0], rule.column_name)


def validate_union_records(record_sets, datasets, rules):
    """Evaluate dataset-scoped rules only against their source rows.

    Job-wide rules still see the complete union, and all outputs retain the
    original row dictionaries and ordering.
    """
    records = [row for rows in record_sets for row in rows]
    failures = {}
    global_rules = [rule for rule in rules if rule.dataset_id is None]
    if global_rules:
        _, global_failures = validate_records(records, global_rules)
        failures.update({item["row_index"]: list(item["failure_reasons"]) for item in global_failures})
    offset = 0
    for dataset, rows in zip(datasets, record_sets):
        scoped_rules = [rule for rule in rules if rule.dataset_id == dataset.id]
        if scoped_rules:
            _, scoped_failures = validate_records(rows, scoped_rules)
            for item in scoped_failures:
                index = offset + item["row_index"]
                reasons = failures.setdefault(index, [])
                reasons.extend(reason for reason in item["failure_reasons"] if reason not in reasons)
        offset += len(rows)
    passed, failed = [], []
    for index, record in enumerate(records):
        reasons = failures.get(index, [])
        if reasons:
            failed.append({"row_index": index, "record": record, "failure_reasons": reasons})
        else:
            passed.append(record)
    return passed, failed


def execute(execution_id: int):
    db = SessionLocal()
    try:
        execution = db.get(Execution, execution_id)
        if execution is None:
            raise ValueError(f"Execution {execution_id} does not exist")
        job = db.get(Job, execution.job_id)
        if job is None:
            raise ValueError(f"Job for execution {execution_id} does not exist")
        execution.status = "running"; execution.started_at = datetime.now(timezone.utc)
        execution.logs = ["Execution started"]
        db.commit()
        destination=job.destination or {"type":"local"}
        # A complete Databricks output configuration in the deployment environment
        # is the configured workspace target for jobs that still carry the legacy
        # local default (the existing UI has no separate output-destination control).
        if (destination.get("type", "local") == "local" and settings.databricks_server_hostname
                and settings.databricks_http_path and settings.databricks_token):
            destination = {**destination, "type": "databricks"}
        datasets = [db.get(Dataset, did) for did in (job.dataset_ids or [])]
        if not datasets:
            raise ValueError("Job must reference at least one dataset")
        if any(d is None for d in datasets): raise ValueError("Job references a missing dataset")
        if len(datasets) > 2 and (job.settings or {}).get("combine_mode") != "union": raise ValueError("A join can combine up to two datasets")
        local = LocalStorage()
        record_sets = [parse_data(d.filename, local.read(d.path)) for d in datasets]
        # Refresh persisted schema metadata from the selected source bytes so a
        # previous import's stale/synthetic columns do not survive execution.
        for dataset, rows in zip(datasets, record_sets):
            if rows:
                dataset.schema_json = schema_from_records(rows)
        if len(record_sets) == 1:
            records = record_sets[0]
        elif (job.settings or {}).get("combine_mode") == "union":
            records = [row for rows in record_sets for row in rows]
        else:
            cfg = job.join_config
            if not cfg: raise ValueError("Jobs with multiple datasets require join_config")
            left_id, right_id = str(cfg.get("left_dataset_id")), str(cfg.get("right_dataset_id"))
            indexes = {str(d.id): rows for d, rows in zip(datasets, record_sets)}
            if left_id not in indexes or right_id not in indexes: raise ValueError("join_config must reference datasets selected by this job")
            left_key, right_key = cfg.get("left_key"), cfg.get("right_key")
            if not left_key or not right_key: raise ValueError("join_config requires left_key and right_key")
            right_index = {}
            for row in indexes[right_id]: right_index.setdefault(str(row.get(right_key)), []).append(row)
            records = []
            for left in indexes[left_id]:
                for right in right_index.get(str(left.get(left_key)), []):
                    merged = {**left, **{(f"{cfg.get('right_prefix', 'right_')}{k}" if k in left else k): v for k,v in right.items()}}
                    records.append(merged)
        selected = job.selected_columns or {}
        if selected:
            if (job.settings or {}).get("combine_mode") == "union" and len(record_sets) > 1:
                # Apply each source's selection before combining rows. A single
                # global column set would leak same-named, unselected fields
                # from another source into this source's union rows.
                selected_sets = []
                for dataset, rows in zip(datasets, record_sets):
                    source_columns = selected.get(str(dataset.id), selected.get(dataset.id,
                        [] if selected else list((dataset.schema_json or {}).get("columns", {}))))
                    selected_sets.append([{key: value for key, value in row.items() if key in source_columns}
                                          for row in rows])
                record_sets = selected_sets
                records = [row for rows in record_sets for row in rows]
            else:
                chosen=set()
                seen_columns=set()
                cfg=job.join_config or {}
                prefix=cfg.get("right_prefix","right_")
                by_id={d.id:d for d in datasets}
                is_join = len(datasets) == 2 and (job.settings or {}).get("combine_mode") != "union"
                if is_join:
                    # The join configuration was validated above; use its dataset order
                    # only for a join. A union has no left/right dataset keys.
                    ordered_datasets=[by_id[int(cfg[k])] for k in ("left_dataset_id","right_dataset_id")]
                else:
                    ordered_datasets=datasets
                for index,dataset in enumerate(ordered_datasets):
                    for col in selected.get(str(dataset.id), selected.get(dataset.id, [])):
                        output_col=f"{prefix}{col}" if is_join and index > 0 and col in seen_columns else col
                        chosen.add(output_col)
                    seen_columns.update((dataset.schema_json or {}).get("columns",{}).keys())
                records = [{k:v for k,v in row.items() if k in chosen} for row in records]
        if execution.rule_snapshot is not None:
            rules = [SimpleNamespace(**item) for item in execution.rule_snapshot]
        else:  # Backwards compatibility for executions queued before migration.
            scope = [Rule.dataset_id.is_(None)]
            if job.dataset_ids:
                scope.append(Rule.dataset_id.in_(job.dataset_ids))
            rules = list(db.scalars(select(Rule).where(Rule.job_id == job.id, or_(*scope))))
        # A rule is valid only for the dataset it names and the actual input
        # columns selected for this job. Never interpret an absent field as
        # null values on every row.
        # One check may have been suggested more than once with different
        # explanatory metadata. Count and execute only distinct enabled,
        # approved rules; pending and disabled checks must never validate data.
        selected_columns = job.selected_columns or {}
        selected_for = lambda dataset: selected_columns.get(str(dataset.id), [] if selected_columns else list((dataset.schema_json or {}).get("columns", {})))
        active_rules = [r for r in rules if r.enabled and r.approved and
                        (r.column_name in selected_for(next((d for d in datasets if d.id == r.dataset_id), datasets[0])) if r.dataset_id is not None
                         else any(r.column_name in selected_for(dataset) for dataset in datasets))]
        validate_rule_bindings(active_rules, job.dataset_ids,
                               {dataset.id: rows for dataset, rows in zip(datasets, record_sets)}, records)
        signatures = {rule_signature(r) for r in active_rules}
        executed_count = len(signatures)
        pending_count = sum(not bool(r.approved) for r in rules)
        disabled_count = sum(not bool(r.enabled) for r in rules)
        skipped_count = sum(not (r.enabled and r.approved) for r in rules)
        if executed_count:
            if (job.settings or {}).get("combine_mode") == "union" and len(record_sets) > 1:
                passed, failed = validate_union_records(record_sets, datasets, active_rules)
            else:
                passed, failed = validate_records(records, active_rules)
        else:
            # No executed rules means no record has been evaluated. Do not
            # classify the untouched input rows as passing.
            passed, failed = [], []
        validation_complete = executed_count > 0
        base = f"outputs/job-{job.id}/execution-{execution.id}"
        dest = LocalStorage()
        passed_path = dest.save(base + "/passed.csv", encode_csv(passed))
        failed_path = dest.save(base + "/failed.json", json.dumps(failed, ensure_ascii=False, indent=2).encode())
        report = {"execution_id": execution.id, "job_id": job.id,
            "status": "succeeded" if validation_complete else "incomplete", "validation_complete": validation_complete,
            "total": len(records) if validation_complete else 0,
            "records_loaded": len(records), "records_evaluated": len(records) if validation_complete else 0,
            "passed": len(passed), "failed": len(failed),
            "rules": {"executed": executed_count, "pending": pending_count, "disabled": disabled_count, "skipped": skipped_count},
            "passed_path": passed_path, "failed_path": failed_path,
            "destination": destination.get("type", "local"),
            "storage_status": {"status": "not_started", "pass_status": "not_started", "fail_status": "not_started"},
            "generated_at": datetime.now(timezone.utc).isoformat()}
        if not validation_complete:
            report["message"] = "Validation incomplete: approve and enable at least one applicable rule, then run validation again. Records were not validated."
            report_path = dest.save(base + "/report.json", json.dumps(report, indent=2).encode())
            execution.status = "incomplete"
            execution.passed_count = execution.failed_count = 0
            execution.report_path = report_path
            execution.logs = ["Execution started", f"Loaded {len(records)} records",
                f"Executed 0 approved and enabled rules; skipped {skipped_count} rules ({pending_count} pending, {disabled_count} disabled; categories can overlap)",
                report["message"]]
            execution.finished_at = datetime.now(timezone.utc)
            db.commit()
            return
        output_columns=list(dict.fromkeys(k for row in passed for k in row))
        output_schema={"columns":{c:{"inferred_type":"string","nullable":any(is_missing(r.get(c)) for r in passed)} for c in output_columns},"row_count":len(passed)}
        output_dataset=Dataset(name=f"{job.name} output {execution.id}",filename="passed.csv",path=passed_path,source_type="derived",schema_json=output_schema)
        db.add(output_dataset); db.flush()
        report["output_dataset_id"] = output_dataset.id
        provider=destination.get("type","local")
        secrets={}
        if provider != "local":
            credential_id=destination.get("credential_id")
            credential=db.get(Credential,credential_id) if credential_id else None
            if credential:
                from cryptography.fernet import Fernet
                if not settings.secret_encryption_key: raise ValueError("SECRET_ENCRYPTION_KEY is required for encrypted cloud credentials")
                secrets=json.loads(Fernet(settings.secret_encryption_key.encode()).decrypt(credential.secret_ciphertext.encode()))
                destination={**(credential.config or {}),**destination}
            elif provider == "databricks":
                # Databricks output may use the documented environment configuration;
                # credentials remain in process memory and are never written to the report.
                destination={**destination,
                    "server_hostname": destination.get("server_hostname") or settings.databricks_server_hostname,
                    "http_path": destination.get("http_path") or settings.databricks_http_path,
                    "catalog": destination.get("catalog") or settings.databricks_catalog,
                    "schema": destination.get("schema") or settings.databricks_schema}
                secrets={"access_token": settings.databricks_token} if settings.databricks_token else {}
                if not destination.get("server_hostname") or not destination.get("http_path") or not secrets.get("access_token"):
                    raise ValueError("Databricks destination requires DATABRICKS_SERVER_HOSTNAME, DATABRICKS_HTTP_PATH, and DATABRICKS_TOKEN, or a saved Databricks credential")
            else:
                raise ValueError(f"Destination {provider} requires credential_id")
            compatible={"s3":"s3","azure_blob":"azure_blob","onelake":"onelake","fabric":"fabric","snowflake":"snowflake","databricks":"databricks"}
            if credential and compatible.get(provider) != credential.provider: raise ValueError(f"Credential provider {credential.provider} does not match destination {provider}")
        dataset_label = "_and_".join(d.name for d in datasets)
        storage_run_id = execution.id
        retry_parent = db.get(Execution, execution.retry_of) if execution.retry_of else None
        while retry_parent:
            storage_run_id = retry_parent.id
            retry_parent = db.get(Execution, retry_parent.retry_of) if retry_parent.retry_of else None
        try:
            failed_indices = {item["row_index"] for item in failed}
            passed_row_indices = [index for index in range(len(records)) if index not in failed_indices]
            remote=write_destination(destination,secrets,base,passed,failed,report,encode_csv(passed),
                dataset_name=dataset_label,job_id=job.id,execution_id=execution.id,
                columns=list(dict.fromkeys(k for row in records for k in row)),
                source_dataset_ids=[dataset.id for dataset in datasets], storage_run_id=storage_run_id,
                passed_row_indices=passed_row_indices)
            report["destination_result"]=remote
            report["storage_status"]={"status":"saved","pass_status":remote.get("pass_status","saved"),
                                      "fail_status":remote.get("fail_status","saved")}
        except Exception as storage_error:
            # Validation remains successful when its independent output write fails.
            report["storage_status"]={"status":"failed","pass_status":getattr(storage_error,"pass_status","failed"),
                                      "fail_status":getattr(storage_error,"fail_status","failed"),
                                      "error":type(storage_error).__name__}
            report["destination_result"]={"type":provider,"status":"failed",
                "pass_status":report["storage_status"]["pass_status"],"fail_status":report["storage_status"]["fail_status"],
                "error":type(storage_error).__name__}
            report["validation_status"] = "succeeded"
            report["status"] = "failed"
            logging.getLogger(__name__).error("Execution %s validation succeeded but destination storage failed (%s)",
                                              execution_id,type(storage_error).__name__)
        # Persist the final report locally after the real destination operation has completed.
        report_path = dest.save(base + "/report.json", json.dumps(report, indent=2).encode())
        storage_succeeded = report["storage_status"]["status"] == "saved"
        report["validation_status"] = "succeeded"
        execution.status = "succeeded" if storage_succeeded else "failed"
        execution.passed_count = len(passed); execution.failed_count = len(failed)
        execution.error = "Destination storage failed; see report storage_status" if report["storage_status"]["status"] == "failed" else None
        execution.report_path = report_path; execution.logs = ["Execution started", f"Loaded {len(records)} records",
            f"Executed {executed_count} approved and enabled rules; skipped {pending_count} pending rules; {disabled_count} disabled rules",
            f"Passed {len(passed)}; failed {len(failed)}", "Outputs saved locally",
            f"Destination storage {report['storage_status']['status']}"]
        execution.finished_at = datetime.now(timezone.utc); db.commit()
    except Exception as exc:
        logging.getLogger(__name__).exception("Execution %s failed", execution_id)
        db.rollback()
        execution = db.get(Execution, execution_id)
        if execution:
            execution.status = "failed"
            execution.passed_count = 0
            execution.failed_count = 0
            execution.report_path = None
            if isinstance(exc, ValueError):
                execution.error = f"Configuration error: {exc}"
            elif isinstance(exc, KeyError):
                missing = exc.args[0] if exc.args else "unknown"
                execution.error = f"Execution failed: required field or configuration key {missing!r} is missing. Check the job configuration and input schema."
            else:
                execution.error = f"Execution failed ({type(exc).__name__}). Verify the job configuration and connected provider."
            execution.finished_at = datetime.now(timezone.utc)
            execution.logs = (execution.logs or []) + [f"Execution failed with {type(exc).__name__}; see backend logs for the diagnostic."]
            db.commit()
    finally:
        db.close()
